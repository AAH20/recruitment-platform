"""Data export service for various formats."""

from __future__ import annotations

import csv
import io
import json
import logging
from typing import Any

logger = logging.getLogger(__name__)


class ExportService:
    """Service for exporting data in various formats."""

    async def export_to_csv(
        self, data: list[dict[str, Any]], fields: list[str] | None = None
    ) -> str:
        """Export data to CSV format."""
        if not data:
            return ""
        if fields is None:
            fields = list(data[0].keys())
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(data)
        return output.getvalue()

    async def export_to_json(self, data: Any, pretty: bool = True) -> str:
        """Export data to JSON format."""
        indent = 2 if pretty else None
        return json.dumps(data, indent=indent, default=str)

    async def export_to_xml(
        self, data: list[dict[str, Any]], root_element: str = "records"
    ) -> str:
        """Export data to XML format."""
        from xml.etree.ElementTree import Element, SubElement, tostring

        root = Element(root_element)
        for item in data:
            record = SubElement(root, "record")
            for key, value in item.items():
                field = SubElement(record, str(key))
                field.text = str(value) if value is not None else ""
        return tostring(root, encoding="unicode")

    async def export_to_excel(
        self, data: list[dict[str, Any]], fields: list[str] | None = None
    ) -> bytes:
        """Export data to Excel format."""
        try:
            import openpyxl
        except ImportError:
            raise ImportError("openpyxl is required for Excel export") from None
        wb = openpyxl.Workbook()
        ws = wb.active
        if not data:
            output = io.BytesIO()
            wb.save(output)
            return output.getvalue()
        if fields is None:
            fields = list(data[0].keys())
        for col, field in enumerate(fields, 1):
            ws.cell(row=1, column=col, value=field)
        for row, item in enumerate(data, 2):
            for col, field in enumerate(fields, 1):
                ws.cell(row=row, column=col, value=item.get(field, ""))
        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    async def export_to_pdf(
        self, data: list[dict[str, Any]], title: str = "Export"
    ) -> bytes:
        """Export data to PDF format."""
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import letter
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import (
                Paragraph,
                SimpleDocTemplate,
                Table,
                TableStyle,
            )
        except ImportError:
            raise ImportError("reportlab is required for PDF export") from None
        output = io.BytesIO()
        doc = SimpleDocTemplate(output, pagesize=letter)
        elements = []
        styles = getSampleStyleSheet()
        elements.append(Paragraph(title, styles["Heading1"]))
        if data:
            headers = list(data[0].keys())
            table_data = [headers]
            for item in data:
                table_data.append([str(item.get(h, "")) for h in headers])
            table = Table(table_data)
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTSIZE", (0, 0), (-1, 0), 12),
                        ("BOTTOMPADDING", (0, 0), (-1, 0), 12),
                        ("BACKGROUND", (0, 1), (-1, -1), colors.beige),
                        ("GRID", (0, 0), (-1, -1), 1, colors.black),
                    ]
                )
            )
            elements.append(table)
        doc.build(elements)
        return output.getvalue()


export_service = ExportService()
